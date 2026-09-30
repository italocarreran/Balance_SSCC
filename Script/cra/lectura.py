# -*- coding: utf-8 -*-
"""
Lectura de los archivos de entrada del CRA por letra de columna y fila
de inicio, que es como los define la trazabilidad ("el bloque A2:D del
origen se lleva a FP!U9:X").

Se lee con openpyxl directo (no con pandas.read_excel) porque la fila
de inicio tiene que ser la fila REAL de Excel: pandas, en modo solo
lectura, puede saltarse filas vacias del principio y correr todo.
"""

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from ..nucleo.utiles import ErrorEntrada, normalizar
from .parametros import EXTENSIONES_EXCEL


# ============================================================
# BUSQUEDA DE ARCHIVOS
# ============================================================

def buscar_por_prefijo(carpeta, prefijo, descripcion):
    """
    El unico archivo Excel de `carpeta` cuyo nombre normalizado empieza
    con `prefijo`. None si no hay ninguno.

    Si hay mas de uno, ErrorEntrada con la lista: no se elige por fecha
    (mismo criterio que el SoC del BESS).
    """

    carpeta = Path(carpeta)

    if not carpeta.is_dir():
        return None

    prefijo = normalizar(prefijo)

    candidatos = sorted(
        ruta for ruta in carpeta.iterdir()
        if ruta.is_file()
        and ruta.suffix.lower() in EXTENSIONES_EXCEL
        and not ruta.name.startswith("~$")      # temporales de Excel abierto
        and normalizar(ruta.stem).startswith(prefijo)
    )

    if len(candidatos) > 1:
        lista = "\n".join(f"  - {ruta.name}" for ruta in candidatos)
        raise ErrorEntrada(
            f"Hay mas de un archivo de {descripcion} en {carpeta}:\n"
            f"{lista}\n"
            f"Deja solo el del periodo que se esta calculando."
        )

    return candidatos[0] if candidatos else None


# ============================================================
# LECTURA POR LETRA
# ============================================================

def elegir_hoja(libro, hoja, ruta):
    """
    La hoja a leer. Si `hoja` es None (no esta confirmada), se usa la
    unica que haya; si hay varias, se para y se listan.
    """

    if hoja is not None:
        if hoja not in libro.sheetnames:
            raise ErrorEntrada(
                f"{Path(ruta).name} no tiene la hoja '{hoja}'. "
                f"Hojas: {', '.join(libro.sheetnames)}."
            )
        return libro[hoja]

    if len(libro.sheetnames) == 1:
        return libro[libro.sheetnames[0]]

    raise ErrorEntrada(
        f"{Path(ruta).name} tiene varias hojas "
        f"({', '.join(libro.sheetnames)}) y todavia no esta definido "
        f"cual se lee. Hay que confirmarlo y dejarlo en "
        f"Script/cra/parametros.py (HOJA_ORIGEN_*)."
    )


def leer_columnas(ruta, hoja, fila_inicio, letras):
    """
    DataFrame con una columna por letra (nombrada con la letra), desde
    `fila_inicio` (1-indexada, como en Excel) hasta la ultima fila con
    algun valor en esas columnas.

    Las filas en las que todas las letras pedidas estan vacias se
    descartan: son el final del bloque o renglones en blanco en medio.
    Los valores son los guardados (data_only): si el origen tiene
    formulas, se lee el resultado que dejo Excel al guardar.
    """

    letras = list(dict.fromkeys(letras))
    indices = [column_index_from_string(letra) for letra in letras]
    minimo, maximo = min(indices), max(indices)

    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise ErrorEntrada(
            f"No se pudo abrir {Path(ruta).name}: {error}"
        ) from error

    try:
        ws = elegir_hoja(libro, hoja, ruta)
        filas = []

        for fila in ws.iter_rows(
            min_row=fila_inicio,
            min_col=minimo,
            max_col=maximo,
            values_only=True,
        ):
            fila = list(fila) + [None] * (maximo - minimo + 1 - len(fila))
            valores = [fila[i - minimo] for i in indices]

            if all(_vacio(v) for v in valores):
                continue

            filas.append(valores)
    finally:
        libro.close()

    return pd.DataFrame(filas, columns=letras, dtype=object)


def _vacio(valor):
    return valor is None or (isinstance(valor, str) and not valor.strip())


# ============================================================
# FECHAS Y TEXTOS "A LA EXCEL"
# ============================================================

ORIGEN_SERIAL_EXCEL = dt.datetime(1899, 12, 30)


def a_fecha_hora(valor):
    """
    datetime a partir de lo que puede traer una celda de fecha: un
    datetime/date (lo normal con openpyxl), un time (solo hora) o un
    numero serial de Excel. None si no se puede.
    """

    if isinstance(valor, dt.datetime):
        return valor

    if isinstance(valor, dt.date):
        return dt.datetime(valor.year, valor.month, valor.day)

    if isinstance(valor, dt.time):
        return dt.datetime.combine(ORIGEN_SERIAL_EXCEL.date(), valor)

    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return ORIGEN_SERIAL_EXCEL + dt.timedelta(days=float(valor))

    return None


def dia_excel(valor):
    """DIA() de Excel. None si la celda no es una fecha."""

    fecha = a_fecha_hora(valor)

    return None if fecha is None else fecha.day


def texto_excel(valor):
    """
    Como queda un valor al concatenarlo con & en Excel: 5 -> "5" (no
    "5.0"), 5.5 -> "5.5", texto tal cual, vacio -> "".
    """

    if _vacio(valor):
        return ""

    if isinstance(valor, bool):
        return "VERDADERO" if valor else "FALSO"

    if isinstance(valor, (int, float)):
        numero = float(valor)
        if numero.is_integer():
            return str(int(numero))
        return repr(numero)

    return str(valor).strip()


def numero_o_nada(valor):
    """float si la celda es numerica (o texto numerico); None si no."""

    if _vacio(valor) or isinstance(valor, bool):
        return None

    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


def entero_si_se_puede(numero):
    """5.0 -> 5; 5.5 -> 5.5; None -> None."""

    if numero is None:
        return None

    return int(numero) if float(numero).is_integer() else numero
