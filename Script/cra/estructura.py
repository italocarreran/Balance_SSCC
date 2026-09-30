# -*- coding: utf-8 -*-
"""
El arbol del caso CRA que dibuja la ventana (Balance_CRA.py).

Una fila por carpeta, archivo u hoja, con su nivel (0 = raiz del caso),
su id (la ventana decide con el que boton le cuelga), su estado ("ok" /
"falta" / "pendiente") y un detalle. No sabe nada de botones.
"""

from openpyxl import load_workbook

from ..nucleo.utiles import ErrorEntrada
from . import parametros as p
from .proceso import SECCIONES
from .rutas import ENTRADAS, buscar_entrada, resolver_rutas

# Hojas del libro original que todavia no tienen su carga definida
# (traz. 3.1 y 17): se muestran para que se vea lo que falta.
HOJAS_PENDIENTES = [
    ("TC", "origen CMg: pendiente"),
    ("COTAS", "origen/formato pendiente"),
    ("RENDIMIENTOS", "maestro: falta definir de donde se lee"),
    ("CONDICION_EMBALSE", "regla/origen pendiente"),
    ("EMPRESAS y unidades candidatas", "maestros: falta definir de donde se leen"),
    ("CÁLCULO_CRA", "calculo: espera las entradas"),
    ("RESUMEN", "calculo: espera CÁLCULO_CRA"),
]


def _fila(nivel, id_, texto, estado, detalle="", ruta=None, es_carpeta=False):
    return {
        "nivel": nivel,
        "id": id_,
        "texto": texto,
        "estado": estado,
        "detalle": detalle,
        "ruta": ruta,
        "es_carpeta": es_carpeta,
    }


def revisar_estructura(carpeta_base, aamm):
    """Lista plana de filas, en el orden en que se dibujan."""

    rutas = resolver_rutas(carpeta_base)
    filas = []

    carpetas = {}
    for id_, carpeta, patron, _ in ENTRADAS:
        carpetas.setdefault(carpeta, []).append((id_, patron))

    for carpeta, entradas in carpetas.items():
        ruta_carpeta = rutas[carpeta]
        existe = ruta_carpeta.is_dir()
        filas.append(_fila(
            0, f"carpeta_{carpeta}", f"{ruta_carpeta.name}/",
            "ok" if existe else "falta", ruta=ruta_carpeta, es_carpeta=True,
        ))

        for id_, patron in entradas:
            patron = patron.replace("<AAMM>", aamm or "<AAMM>")
            try:
                ruta = buscar_entrada(rutas, id_, aamm or "")
            except ErrorEntrada as error:
                filas.append(_fila(
                    1, id_, patron, "falta",
                    str(error).splitlines()[0], ruta=ruta_carpeta,
                ))
                continue

            if ruta is None:
                filas.append(_fila(
                    1, id_, patron, "falta", "no esta", ruta=ruta_carpeta,
                ))
            else:
                filas.append(_fila(1, id_, ruta.name, "ok", ruta=ruta))

    salida = rutas["salida"]
    hojas_escritas = _hojas_de(salida)
    filas.append(_fila(
        0, "salida", salida.name,
        "ok" if salida.exists() else "pendiente",
        "" if salida.exists() else "todavia no se genero",
        ruta=salida,
    ))

    por_hoja = {hoja: seccion for seccion, (hoja, _) in SECCIONES.items()}
    for hoja in p.ORDEN_HOJAS_SALIDA:
        escrita = hoja in hojas_escritas
        detalle = "" if escrita else "sin generar"
        if hoja == p.HOJA_PRORRATA:
            detalle = "solo el bloque 1 (matriz de prorratas)"
        filas.append(_fila(
            1, f"hoja_{por_hoja[hoja]}", f"hoja '{hoja}'",
            "ok" if escrita else "pendiente", detalle, ruta=salida,
        ))

    for hoja, motivo in HOJAS_PENDIENTES:
        filas.append(_fila(1, "sin_definir", f"hoja '{hoja}'", "pendiente", motivo))

    return filas


def _hojas_de(ruta):
    if not ruta.exists():
        return set()

    try:
        libro = load_workbook(ruta, read_only=True)
    except Exception:
        return set()

    try:
        return set(libro.sheetnames)
    finally:
        libro.close()
