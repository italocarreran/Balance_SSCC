# -*- coding: utf-8 -*-
"""
Escritura de Balance_CRA.xlsx: se reemplazan solo las hojas que se
regeneran y el resto del libro se conserva tal cual estaba (mismo
criterio que Balance_BESS.xlsx).
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

from ..nucleo.formato import formatear_hoja
from ..nucleo.utiles import ErrorEntrada
from .parametros import ORDEN_HOJAS_SALIDA


def escribir_hojas(ruta_salida, hojas):
    """
    hojas: {nombre de hoja: DataFrame}. Reemplaza esas hojas en el
    libro (lo crea si no existe), deja las demas como estaban, ordena
    segun ORDEN_HOJAS_SALIDA (las que no esten ahi van al final) y
    formatea solo las que se escribieron.
    """

    ruta_salida = Path(ruta_salida)

    if ruta_salida.exists():
        try:
            libro = load_workbook(ruta_salida)
        except Exception as error:
            raise ErrorEntrada(
                f"No se pudo abrir {ruta_salida.name} para actualizarlo "
                f"(¿esta abierto en Excel?): {error}"
            ) from error
    else:
        libro = Workbook()
        libro.remove(libro.active)

    for nombre, df in hojas.items():
        if nombre in libro.sheetnames:
            libro.remove(libro[nombre])

        ws = libro.create_sheet(nombre)
        ws.append([str(c) for c in df.columns])

        for fila in df.itertuples(index=False):
            ws.append([_celda(v) for v in fila])

        formatear_hoja(ws)

    orden = {nombre: i for i, nombre in enumerate(ORDEN_HOJAS_SALIDA)}
    libro._sheets.sort(key=lambda ws: orden.get(ws.title, len(orden)))
    libro.active = 0

    try:
        libro.save(ruta_salida)
    except PermissionError as error:
        raise ErrorEntrada(
            f"No se pudo guardar {ruta_salida.name}: esta abierto en "
            f"Excel. Cerralo y volve a intentar."
        ) from error


def _celda(valor):
    if valor is None:
        return None

    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass

    if hasattr(valor, "item"):         # numpy -> python
        return valor.item()

    return valor
