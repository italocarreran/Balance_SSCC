# -*- coding: utf-8 -*-
"""
Auxiliares/centrales_cra.xlsx: los maestros del CRA, en un solo libro.

    hoja "centrales_cra"  Configuracion                 -> filtra cvar_cra
    hoja "empresas"       UNIDAD/CONFIGURACION, EMPRESA -> (CÁLCULO_CRA)
    hoja "diccionario"    UNIDAD/CONFIGURACION, FD
                          -> que unidades del SSCC_Desempeño entran a
                             las hojas FD_* (la misma lista para CPF, CSF
                             y CTF: el usuario dejo una sola columna
                             porque eran iguales)

Cada hoja trae un titulo ("Cuadro N° ...") arriba de los encabezados: la
fila de encabezados se busca por su texto, nunca por posicion (mismo
criterio que Resumen BESS en el BESS). Los datos siguen hasta la primera
fila con la primera columna vacia.

En el diccionario una celda puede listar varias unidades separadas por
";" (ej. "HE CIPRESES U1; HE CIPRESES U2; HE CIPRESES U3").
"""

import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from ..nucleo.utiles import ErrorEntrada, normalizar

HOJA_CONFIGURACIONES = "centrales_cra"
HOJA_EMPRESAS = "empresas"
HOJA_DICCIONARIO = "diccionario"

COL_CONFIGURACION = "Configuracion"
COL_UNIDAD = "UNIDAD/CONFIGURACION"
COL_EMPRESA = "EMPRESA"
COL_FD = "FD"

SEPARADOR_UNIDADES = ";"

# Hasta que fila se busca el encabezado (el titulo va arriba).
FILAS_BUSQUEDA_ENCABEZADO = 30


def leer_tabla(ruta, hoja, columnas):
    """
    Lista de dicts {columna: valor} de `hoja`, buscando la fila que
    tiene TODOS los textos de `columnas`. Termina en la primera fila
    con la primera de `columnas` vacia.
    """

    ruta = Path(ruta)

    try:
        filas = _filas_xml(ruta, hoja)
    except ErrorEntrada:
        raise
    except Exception:
        filas = None            # formato raro: se lee con openpyxl

    if filas is not None:
        tabla = _tabla_desde_filas(iter(filas), columnas, ruta, hoja)
    else:
        tabla = _tabla_openpyxl(ruta, hoja, columnas)

    if not tabla:
        raise ErrorEntrada(f"{ruta.name}, hoja '{hoja}': la tabla esta vacia.")

    return tabla


def _tabla_openpyxl(ruta, hoja, columnas):
    try:
        libro = load_workbook(ruta, read_only=True, data_only=True)
    except Exception as error:
        raise ErrorEntrada(f"No se pudo abrir {ruta.name}: {error}") from error

    try:
        nombres = {normalizar(n): n for n in libro.sheetnames}
        if normalizar(hoja) not in nombres:
            raise ErrorEntrada(
                f"{ruta.name} no tiene la hoja '{hoja}'. Hojas: "
                f"{', '.join(libro.sheetnames)}."
            )
        filas = libro[nombres[normalizar(hoja)]].iter_rows(values_only=True)
        return _tabla_desde_filas(filas, columnas, ruta, hoja)
    finally:
        libro.close()


# ============================================================
# LECTURA DIRECTA DEL XML
#
# El centrales_cra.xlsx real trae un styles.xml de ~11 MB (estilos
# arrastrados del libro de origen) y openpyxl lo procesa entero al abrir:
# ~9 s por apertura para tres tablas chicas. Los maestros son solo
# valores, asi que se leen directo del zip (workbook -> hoja -> celdas,
# con los textos compartidos) sin tocar los estilos. Si el archivo trae
# algo que este lector no entiende, se vuelve a openpyxl.
# ============================================================

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_NS_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CELDA = re.compile(r"([A-Z]+)(\d+)")


def _filas_xml(ruta, hoja):
    """Lista de tuplas (una por fila de Excel, desde la 1) con los valores."""

    try:
        zf = zipfile.ZipFile(ruta)
    except (OSError, zipfile.BadZipFile) as error:
        raise ErrorEntrada(f"No se pudo abrir {ruta.name}: {error}") from error

    with zf:
        libro = ET.fromstring(zf.read("xl/workbook.xml"))
        hojas = {
            h.get("name"): h.get(f"{_NS_REL}id")
            for h in libro.iter(f"{_NS}sheet")
        }
        nombres = {normalizar(n): n for n in hojas}
        if normalizar(hoja) not in nombres:
            raise ErrorEntrada(
                f"{ruta.name} no tiene la hoja '{hoja}'. Hojas: "
                f"{', '.join(hojas)}."
            )

        relaciones = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        destino = {
            r.get("Id"): r.get("Target")
            for r in relaciones.iter(f"{_NS_PKG}Relationship")
        }[hojas[nombres[normalizar(hoja)]]]
        destino = (destino.lstrip("/") if destino.startswith("/")
                   else posixpath.normpath(posixpath.join("xl", destino)))

        compartidos = []
        if "xl/sharedStrings.xml" in zf.namelist():
            for si in ET.fromstring(zf.read("xl/sharedStrings.xml")).iter(f"{_NS}si"):
                compartidos.append("".join(
                    t.text or "" for t in si.iter(f"{_NS}t")
                ))

        filas = {}
        for c in ET.fromstring(zf.read(destino)).iter(f"{_NS}c"):
            letra, numero = _CELDA.fullmatch(c.get("r")).groups()
            valor = _valor_celda(c, compartidos)
            if valor is not None:
                filas.setdefault(int(numero), {})[
                    column_index_from_string(letra) - 1] = valor

    ultima = max(filas, default=0)
    return [
        tuple(
            filas.get(n, {}).get(j)
            for j in range(max(filas.get(n, {0: None}), default=0) + 1)
        )
        for n in range(1, ultima + 1)
    ]


def _valor_celda(c, compartidos):
    tipo = c.get("t")

    if tipo == "inlineStr":
        return "".join(t.text or "" for t in c.iter(f"{_NS}t"))

    v = c.find(f"{_NS}v")
    if v is None or v.text is None:
        return None

    if tipo == "s":
        return compartidos[int(v.text)]
    if tipo in ("str", "e"):
        return v.text
    if tipo == "b":
        return v.text == "1"

    numero = float(v.text)
    return int(numero) if numero.is_integer() else numero


def _tabla_desde_filas(filas, columnas, ruta, hoja):
    buscadas = [normalizar(c) for c in columnas]
    posiciones = None

    for i, fila in enumerate(filas):
        textos = [normalizar(v) for v in fila]
        if all(b in textos for b in buscadas):
            posiciones = [textos.index(b) for b in buscadas]
            break
        if i + 1 >= FILAS_BUSQUEDA_ENCABEZADO:
            break

    if posiciones is None:
        raise ErrorEntrada(
            f"{Path(ruta).name}, hoja '{hoja}': no se encontro la fila de "
            f"encabezados ({', '.join(columnas)})."
        )

    tabla = []
    for fila in filas:
        valores = [fila[j] if j < len(fila) else None for j in posiciones]
        if valores[0] is None or str(valores[0]).strip() == "":
            break
        tabla.append({
            col: (v.strip() if isinstance(v, str) else v)
            for col, v in zip(columnas, valores)
        })

    return tabla


def leer_configuraciones(ruta):
    """Las configuraciones del CRA (hoja centrales_cra), en orden."""

    return [
        fila[COL_CONFIGURACION]
        for fila in leer_tabla(ruta, HOJA_CONFIGURACIONES, [COL_CONFIGURACION])
    ]


def leer_empresas(ruta):
    """{unidad/configuracion: empresa} (hoja empresas)."""

    return {
        fila[COL_UNIDAD]: fila[COL_EMPRESA]
        for fila in leer_tabla(ruta, HOJA_EMPRESAS, [COL_UNIDAD, COL_EMPRESA])
    }


def leer_unidades_fd(ruta):
    """
    Las unidades del SSCC_Desempeño que entran a las hojas FD_* (columna
    FD del diccionario): separando las listas por ";" y sin repetir, en
    el orden del maestro. Es la misma lista para CPF, CSF y CTF.
    """

    tabla = leer_tabla(ruta, HOJA_DICCIONARIO, [COL_UNIDAD, COL_FD])

    unidades = []
    for fila in tabla:
        celda = fila[COL_FD]
        if celda is None:
            continue
        for unidad in str(celda).split(SEPARADOR_UNIDADES):
            unidad = unidad.strip()
            if unidad and unidad not in unidades:
                unidades.append(unidad)

    return unidades
