# -*- coding: utf-8 -*-
"""CMg en los meses de cambio de hora, y la FECHA AAAAMMDD sin miles.

El CMg se cruza con Medidores por barra + "Cuarto de Hora" del mes. El
CSV trae la hora que no existe (1 del 6-09-2026), en cero: si se queda,
desde ahi el CMg queda 4 cuartos adelante de Medidores (2876 cuartos en
septiembre) y todo el resto del mes cruza corrido una hora.
"""

import datetime
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from Script.Cmg import Extrae_CMG_barras as cmg
from Script.nucleo.cambio_hora import cambios_del_periodo
from Script.nucleo.formato import FORMATO_CLAVE, FORMATO_ENTERO, formatear_hoja


def _csv(carpeta, anio, mes, horas_por_dia):
    """horas_por_dia: {dia: lista de HORA (periodos)}; el resto 1..24."""

    filas = []
    dias = pd.Timestamp(anio, mes, 1).days_in_month
    for dia in range(1, dias + 1):
        for hora in horas_por_dia.get(dia, range(1, 25)):
            for minuto in (0, 15, 30, 45):
                valor = "0" if (dia, hora) in horas_por_dia.get("cero", ()) \
                    else f"{dia}{hora:02d},{minuto}"
                filas.append({
                    "FECHA": f"{anio}{mes:02d}{dia:02d}", "HORA": hora,
                    "MINUTO": minuto, "BARRA": "BARRA_X", "X": "",
                    "CMg[CLP/KWh]": valor, "CMg[USD/MWh]": "",
                })
    ruta = Path(carpeta) / "cmg.csv"
    pd.DataFrame(filas).to_csv(ruta, sep=";", index=False, encoding="latin1")
    return ruta


def _inexistentes(aamm):
    return [
        (c.fecha, p) for c in cambios_del_periodo(aamm)
        for p in c.periodos_inexistentes
    ]


class TestCmgCambioDeHora(unittest.TestCase):
    def test_septiembre_saca_la_hora_que_no_existe(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = _csv(carpeta, 2026, 9, {"cero": {(6, 1)}})
            log = []
            df, resumen = cmg.construir_cmg_desde_csv(
                ruta, ["BARRA_X"], registrar=log.append,
                horas_inexistentes=_inexistentes("2609"),
            )
        # mismo total que Medidores: 30 dias x 96 - 4
        self.assertEqual(int(df[cmg.COLUMNA_CUARTO].max()), 2876)
        dia6 = df[df["FECHA"] == 20260906]
        self.assertNotIn(1, set(dia6["HORA"]))
        # la 01:00 del 6-09 es el cuarto 481, igual que en Medidores
        primero = dia6.sort_values(cmg.COLUMNA_CUARTO).iloc[0]
        self.assertEqual((int(primero["HORA"]), int(primero[cmg.COLUMNA_CUARTO])),
                         (2, 5 * 96 + 1))
        # el 7-09 sigue calzando (antes quedaba corrido 4 cuartos)
        dia7 = df[df["FECHA"] == 20260907].sort_values(cmg.COLUMNA_CUARTO)
        self.assertEqual(int(dia7[cmg.COLUMNA_CUARTO].iloc[0]), 6 * 96 - 4 + 1)
        self.assertTrue(any("venian todas en cero" in l for l in log), log)

    def test_sin_el_aviso_quedaria_corrido(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = _csv(carpeta, 2026, 9, {"cero": {(6, 1)}})
            df, _ = cmg.construir_cmg_desde_csv(
                ruta, ["BARRA_X"], registrar=lambda *_: None
            )
        self.assertEqual(int(df[cmg.COLUMNA_CUARTO].max()), 2880)

    def test_avisa_si_la_hora_que_se_saca_traia_valores(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = _csv(carpeta, 2026, 9, {})
            log = []
            cmg.construir_cmg_desde_csv(
                ruta, ["BARRA_X"], registrar=log.append,
                horas_inexistentes=_inexistentes("2609"),
            )
        self.assertTrue(any("OJO" in l for l in log), log)

    def test_abril_con_hora_25_no_cambia(self):
        with tempfile.TemporaryDirectory() as carpeta:
            ruta = _csv(carpeta, 2026, 4, {4: range(1, 26)})
            df, _ = cmg.construir_cmg_desde_csv(
                ruta, ["BARRA_X"], registrar=lambda *_: None,
                horas_inexistentes=_inexistentes("2604"),
            )
        self.assertEqual(int(df[cmg.COLUMNA_CUARTO].max()), 30 * 96 + 4)
        dia4 = df[df["FECHA"] == 20260404].sort_values(cmg.COLUMNA_CUARTO)
        self.assertEqual(dia4["HORA"].tolist()[-8:], [24] * 4 + [25] * 4)


class TestFormatoClaves(unittest.TestCase):
    def test_fecha_aaaammdd_y_anio_sin_separador_de_miles(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["FECHA", "Año", "Cuarto de Hora"])
        for i in range(3):
            ws.append([20260905 + i, 2026, 1500 + i])
        formatear_hoja(ws)
        self.assertEqual(ws["A2"].number_format, FORMATO_CLAVE)
        self.assertEqual(ws["B2"].number_format, FORMATO_CLAVE)
        self.assertEqual(ws["C2"].number_format, FORMATO_ENTERO)


if __name__ == "__main__":
    unittest.main()
