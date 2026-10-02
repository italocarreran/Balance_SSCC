# -*- coding: utf-8 -*-
"""
fp_<AAMM> y cvar_cra_<AAMM> desde las politicas de operacion
(Script/Politicas/Costos_Variables.py y Script/cra/politicas.py), con un
arbol de red falso armado en un temporal.
"""

import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import Workbook

from Script import cra
from Script.Politicas import Costos_Variables as cv
from Script.cra import parametros as p

AAMM = "2602"          # febrero 2026: 28 dias
DIAS = 28


def _politica_csv(ruta, dia, valor_a=10.0, valor_b=20.0, horas=range(1, 25)):
    """PO*.csv: primera columna = configuracion (encabezado = fecha)."""

    filas = {f"{dia:02d}-02-2026": ["CONF_A", "CONF_B", "CONF_X"]}
    for hora in horas:
        filas[str(hora)] = [
            valor_a + hora,
            # CONF_B sin dato en la hora 5: la celda vacia se descarta.
            None if hora == 5 else valor_b,
            1.0,
        ]
    ruta.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(filas).to_csv(ruta, index=False)


def _politica_xlsx(ruta):
    """PO*.xlsx: los FP en la TERCERA hoja, encabezado en la fila 2, B:Z."""

    libro = Workbook()
    libro.active.title = "Hoja A"
    libro.create_sheet("Hoja B")
    ws = libro.create_sheet("Factores")
    ws["B1"] = "titulo"
    ws.cell(row=2, column=2, value="BarNom        ")
    for hora in range(1, 25):
        ws.cell(row=2, column=2 + hora, value=hora)
    ws.cell(row=3, column=2, value="Quillota 220")
    for hora in range(1, 25):
        ws.cell(row=3, column=2 + hora, value=1 + hora / 100)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)


class PruebaPoliticas(unittest.TestCase):

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        raiz = Path(self._tmp.name)
        self.politicas = raiz / "progdiar_SEN"
        self.pid = raiz / "PID"
        self.prg = raiz / "CMgReales"

        for dia in range(1, DIAS + 1):
            _politica_csv(self.politicas / "26" / f"PO{AAMM}{dia:02d}.csv", dia)
            _politica_xlsx(self.politicas / "26" / f"PO{AAMM}{dia:02d}.xlsx")

        # Dia 3: reprogramacion desde la hora 10 (CONF_A pasa a 500+h).
        (self.prg / AAMM / "Politicas").mkdir(parents=True)
        (self.prg / AAMM / "Politicas" / f"PRG{AAMM}03_10.xlsx").touch()
        _politica_csv(
            cv.ruta_politica_pid(AAMM, 3, 10, self.pid), 3, valor_a=500.0,
        )
        # Dia 4: hay programa de la hora 15 pero no su politica.
        (self.prg / AAMM / "Politicas" / f"PRG{AAMM}04_15.xlsx").touch()

        self.base = raiz / "CRA 2602"
        cra.crear_carpetas_caso(self.base)
        pd.DataFrame({"Configuracion": ["CONF_A", "CONF_B", "CONF_NUEVA"],
                      "Central": ["A", "B", "N"]}).to_excel(
            self.base / p.CARPETA_AUXILIARES / p.ARCHIVO_CENTRALES_CRA,
            index=False,
        )

        self._parches = [
            mock.patch.object(cv, "RAIZ_POLITICAS", str(self.politicas)),
            mock.patch.object(cv, "RAIZ_PID", str(self.pid)),
            mock.patch.object(cv, "RAIZ_PRG", str(self.prg)),
        ]
        for parche in self._parches:
            parche.start()

    def tearDown(self):
        for parche in self._parches:
            parche.stop()
        self._tmp.cleanup()

    def test_rutas_como_entradas_sscc(self):
        self.assertEqual(
            cv.ruta_politica(AAMM, 3, "csv", "R"), Path("R/26/PO260203.csv"))
        self.assertEqual(
            cv.ruta_politica_pid(AAMM, 3, 7, "P"),
            Path("P/2026/PID_202602/PID_20260203/Publicacion/PID_CDC_07/"
                 "PO260203_07.csv"))
        self.assertEqual(
            cv.ruta_programa_pid(AAMM, 3, 7, "G"),
            Path("G/2602/Politicas/PRG260203_07.xlsx"))
        self.assertEqual(cv.nombre_salida("fp", AAMM), "fp_2602_1_28.xlsx")

    def test_cvar_filtra_centrales_cra_y_aplica_pid(self):
        avisos = []
        cvar = cv.construir(
            AAMM, que=["cvar"],
            ruta_centrales_cra=(self.base / p.CARPETA_AUXILIARES
                                / p.ARCHIVO_CENTRALES_CRA),
            registrar=avisos.append,
        )["cvar"]

        self.assertEqual(list(cvar.columns)[:4], cv.COLUMNAS_CVAR)
        self.assertEqual(set(cvar["Configuracion"]), {"CONF_A", "CONF_B"})
        # CONF_A 24 h x 28 dias; CONF_B sin la hora 5.
        self.assertEqual(len(cvar), DIAS * 24 + DIAS * 23)

        a = cvar[cvar["Configuracion"] == "CONF_A"].set_index(["dia", "hora"])
        self.assertEqual(a.loc[(3, 9), "Cvar"], 19.0)      # antes del PID
        self.assertEqual(a.loc[(3, 10), "Cvar"], 510.0)    # desde el PID
        self.assertEqual(a.loc[(4, 20), "Cvar"], 30.0)     # PID sin politica

        self.assertTrue(any("hora 15" in m and "no su politica" in m
                            for m in avisos))
        self.assertTrue(any("CONF_NUEVA" in m for m in avisos))

    def test_fp_tercera_hoja(self):
        fp = cv.construir(AAMM, que=["fp"], registrar=lambda _m: None)["fp"]

        self.assertEqual(list(fp.columns), cv.COLUMNAS_FP)
        self.assertEqual(len(fp), DIAS * 24)
        primera = fp.iloc[0]
        self.assertEqual(primera["BarNom"], "Quillota 220")
        self.assertEqual(primera["Hora"], 1)
        self.assertAlmostEqual(primera["FP"], 1.01)

    def test_falta_un_dia_se_detiene(self):
        (self.politicas / "26" / f"PO{AAMM}15.csv").unlink()

        with self.assertRaises(cv.ErrorPoliticas) as ctx:
            cv.construir(AAMM, que=["cvar"], ruta_centrales_cra=(
                self.base / p.CARPETA_AUXILIARES / p.ARCHIVO_CENTRALES_CRA))
        self.assertIn(f"PO{AAMM}15.csv", str(ctx.exception))

    def test_generar_y_leer_las_hojas_del_caso(self):
        escritos = cra.generar_politicas(self.base, AAMM)

        self.assertEqual(
            sorted(r.name for r in escritos),
            ["cvar_cra_2602_1_28.xlsx", "fp_2602_1_28.xlsx"])

        # Las hojas FP y CO del Balance_CRA los leen tal cual.
        rutas = cra.resolver_rutas(self.base)
        fp = cra.construir_fp(cra.buscar_entrada(rutas, "fp", AAMM))
        co = cra.construir_co(cra.buscar_entrada(rutas, "co", AAMM))
        self.assertEqual(len(fp), DIAS * 24)
        self.assertEqual(co.iloc[0].tolist(), ["CONF_A", 1, 1, 11])

    def test_no_deja_dos_archivos_del_periodo(self):
        pd.DataFrame({"x": [1]}).to_excel(
            self.base / p.CARPETA_FP / "fp_2602_a_mano.xlsx", index=False)

        with self.assertRaises(cra.ErrorEntrada):
            cra.generar_politicas(self.base, AAMM, que=["fp"])


if __name__ == "__main__":
    unittest.main()
