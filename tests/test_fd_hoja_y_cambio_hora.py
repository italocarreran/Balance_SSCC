# -*- coding: utf-8 -*-
"""Hoja FD sin auxiliares, encabezados con color, y FD en el cambio de hora.

- la hoja no trae 'id' ni la copia de 'Hora Mes' (antes A, M, Q, AE);
  'Central' queda al final de cada tabla y CSF y CPF van separadas por
  una columna. Al leerla vuelven las auxiliares: el calculo ve lo mismo;
- encabezados #244062, letra blanca en negrita, y contorno enmarcado;
- Hora Mes de FD correlativa como la de Medidores: el 6-09-2026 la hora
  0 no existe (queda en 0 y sin Hora Mes) y la 1 es la 121; en abril el
  sabado llega a la hora 24 (periodo 25) y el domingo 0 es la 98.
"""

import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from Script import nucleo
from Script.nucleo.formato import COLOR_ENCABEZADO
from Script.nucleo.hojas_entrada import (
    _construir_bloque_fd_cpf, _construir_bloque_fd_csf,
)


def _crudo_csf(horas_por_fecha, valor=0.9):
    filas = [
        (pd.Timestamp(fecha), hora)
        for fecha, horas in horas_por_fecha for hora in horas
    ]
    return pd.DataFrame({
        "D": [f for f, _ in filas], "E": [h for _, h in filas],
        "F": "BESS_A", "G": "No Participó", "H": 1, "I": 1, "J": valor,
    })


def _crudo_cpf(horas_por_fecha):
    filas = [
        (pd.Timestamp(fecha), hora)
        for fecha, horas in horas_por_fecha for hora in horas
    ]
    return pd.DataFrame({
        "T": [f for f, _ in filas], "U": [h for _, h in filas],
        "V": "BESS_A", "W": 1, "X": 1, "Y": 1, "Z": 1, "AA": 0.8, "AB": "Si",
    })


class TestFdCambioDeHora(unittest.TestCase):
    def test_septiembre(self):
        crudo = _crudo_csf([("2026-09-05", range(24)), ("2026-09-06", range(24)),
                            ("2026-09-07", range(2))])
        crudo.loc[(crudo["D"] == "2026-09-06") & (crudo["E"] == 0), "J"] = 0.5
        fd = _construir_bloque_fd_csf(crudo)
        hora_mes = fd.iloc[:, 1]
        dia6 = fd[fd["Fecha"] == pd.Timestamp("2026-09-06")]
        cero = dia6[dia6["Hora"] == 0]
        self.assertTrue(cero.iloc[:, 1].isna().all())
        self.assertEqual(cero["id"].tolist(), [""])
        self.assertEqual(cero["Factor de Desempeño\n (Fd_CSF)"].tolist(), [0])
        self.assertEqual(cero["CSF(+)"].tolist(), [0])
        self.assertEqual(fd.attrs["hora_inexistente"], (1, 1))
        uno = dia6[dia6["Hora"] == 1]
        self.assertEqual(int(uno.iloc[0, 1]), 121)  # = Medidores 06-09 01:00
        self.assertEqual(uno["id"].tolist(), ["121BESS_A"])
        # el 5-09 no cambia y el 7-09 sigue correlativo
        self.assertEqual(int(hora_mes.iloc[23]), 120)
        self.assertEqual(int(fd.iloc[-2, 1]), 6 * 24 - 1 + 1)

    def test_septiembre_cpf(self):
        fd = _construir_bloque_fd_cpf(
            _crudo_cpf([("2026-09-06", range(3))])
        )
        self.assertTrue(pd.isna(fd.iloc[0, 1]))
        self.assertEqual(fd["CPF(+)"].tolist()[0], 0)
        self.assertEqual(fd["Cuenta con equipo\nregistrador validado"].tolist()[0], "Si")
        self.assertEqual([int(v) for v in fd.iloc[1:, 1]], [121, 122])

    def test_abril(self):
        fd = _construir_bloque_fd_csf(
            _crudo_csf([("2026-04-04", range(25)), ("2026-04-05", range(2))])
        )
        self.assertEqual(
            [int(v) for v in fd.iloc[-4:, 1]], [96, 97, 98, 99]
        )

    def test_mes_normal_igual_que_la_formula(self):
        fd = _construir_bloque_fd_csf(_crudo_csf([("2026-07-15", range(24))]))
        self.assertEqual(
            [int(v) for v in fd.iloc[:, 1]],
            [14 * 24 + h + 1 for h in range(24)],
        )


class TestHojaFd(unittest.TestCase):
    def test_sin_auxiliares_con_central_y_color(self):
        csf = _construir_bloque_fd_csf(_crudo_csf([("2026-07-01", range(3))]))
        cpf = _construir_bloque_fd_cpf(_crudo_cpf([("2026-07-01", range(5))]))
        csf["Central"] = "SAE UNO"
        cpf["Central"] = "SAE UNO"

        with tempfile.TemporaryDirectory() as carpeta:
            ruta = Path(carpeta) / "Consolidado.xlsx"
            nucleo.escribir_salida(
                pd.DataFrame({"Mes": [7]}), ruta, [], [],
                df_fd_csf=csf, df_fd_cpf=cpf,
            )
            ws = openpyxl.load_workbook(ruta)["FD"]
            encabezado = [c.value for c in ws[1]]
            leido_csf, leido_cpf = nucleo.leer_fd_consolidado(
                ruta, registrar=lambda _: None
            )

        n_csf = len(nucleo.NOMBRES_FD_CSF) - 2 + 1
        self.assertEqual(encabezado[0], "Hora Mes")
        self.assertNotIn("id", encabezado)
        self.assertEqual(encabezado[n_csf - 1], "Central")
        self.assertIsNone(encabezado[n_csf])       # una columna de separacion
        self.assertEqual(encabezado[n_csf + 1], "Hora Mes")
        self.assertEqual(encabezado[-1], "Central")
        self.assertEqual(encabezado.count("Hora Mes"), 2)  # una por tabla

        celda = ws.cell(1, 1)
        self.assertEqual(celda.fill.start_color.rgb[-6:], COLOR_ENCABEZADO)
        self.assertTrue(celda.font.bold)
        self.assertEqual(celda.font.color.rgb[-6:], "FFFFFF")
        self.assertEqual(ws.cell(2, 1).border.left.style, "thin")
        self.assertIsNone(ws.cell(1, n_csf + 1).fill.fill_type)

        # Al leerla vuelve todo, con 'id' rearmado
        self.assertEqual(
            list(leido_csf.columns),
            list(nucleo.NOMBRES_FD_CSF.values()) + ["Central"],
        )
        self.assertEqual(leido_csf["id"].tolist(), csf["id"].tolist())
        self.assertEqual(leido_cpf["id"].tolist(), cpf["id"].tolist())
        self.assertEqual(len(leido_cpf), 5)


if __name__ == "__main__":
    unittest.main()
