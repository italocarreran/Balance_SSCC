# -*- coding: utf-8 -*-
"""
Hojas de entrada del CRA (Script/cra): ENERGIA, FP, CO y SC y CO, con
archivos sinteticos armados con la forma que describe la trazabilidad
(docs/Trazabilidad_CRA_Periodo_Generico_v3.md, 6.1, 6.4 y 6.6).
"""

import datetime as dt
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string

from Script import cra
from Script.cra import parametros as p


def _libro(ruta, filas_por_celda, hojas_extra=(), hoja="Hoja1"):
    """filas_por_celda: {(fila, letra): valor}."""

    libro = Workbook()
    ws = libro.active
    ws.title = hoja
    for (fila, letra), valor in filas_por_celda.items():
        ws.cell(row=fila, column=column_index_from_string(letra), value=valor)
    for nombre in hojas_extra:
        libro.create_sheet(nombre)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta


def _energia(ruta, anio=2026, mes=8):
    celdas = {(1, "B"): "titulo", (2, "B"): "Central"}
    datos = [
        ("PEHUENCHE_U1", "PM1", dt.datetime(anio, mes, 1, 0, 0), 0, 100, 5),
        ("PEHUENCHE_U1", "PM1", dt.datetime(anio, mes, 1, 0, 45), 0, 110, 0),
        ("RALCO_U1", "PM2", dt.datetime(anio, mes, 2, 23, 30), 23, 90, 1),
    ]
    for i, (u, pm, fecha, hora, d, r) in enumerate(datos, start=3):
        celdas.update({
            (i, "B"): u, (i, "C"): pm, (i, "D"): fecha.year,
            (i, "E"): fecha.month, (i, "F"): fecha.day, (i, "G"): fecha,
            (i, "H"): d, (i, "I"): r, (i, "J"): hora,
        })
    return _libro(ruta, celdas)


def _reporte_cra(ruta):
    """CSV con ";" y coma decimal, fecha como texto dia-mes-año."""

    lineas = [
        ";".join(f"enc {letra}" for letra in "ABCDEFGHIJKLMNO"),
        "202608;CO;05-08-2026 00:00;x;17;x;x;RAPEL-1;1234,5;"
        "0,25;0,25;0,5;0;0;0",
        "202608;CO;05-08-2026 00:00;x;18;x;x;PE-TOLPANSUR;1;1;0;0;0;0;0",
        "202608;CO;05-08-2026 00:00;x;19;x;x;NUEVA-7;1;1;0;0;0;0;0",
    ]
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text("\n".join(lineas) + "\n", encoding="latin-1")
    return ruta


def _sscc_desempeno(ruta):
    libro = Workbook()
    libro.remove(libro.active)
    for hoja, ultima in (("CPF Horario", "J"), ("CSF Horario", "H"),
                         ("CTF Horario", "I")):
        ws = libro.create_sheet(hoja)
        ws["B1"] = "titulo"
        letras = "BCDEFGHIJ"[: "BCDEFGHIJ".index(ultima) + 1]
        for letra in letras:
            ws[f"{letra}11"] = f"Col {letra}"
        ws["B11"], ws["C11"], ws["D11"] = "Fecha", "Hora", "Unidad"
        ws["B12"], ws["C12"], ws["D12"] = dt.datetime(2026, 8, 1), 0, "RAPEL_U1"
        ws[f"{ultima}12"] = 0.9
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta


def _prorrata(ruta, periodos=31 * 96):
    libro = Workbook()
    ws = libro.active
    ws.title = "Prorrata 15min"
    ws.append(["Cuarto de Hora", "Suministrador", "Prorrata"])
    for cuarto in range(1, periodos + 1):
        ws.append([cuarto, "ENGIE", 0.75])
        ws.append([cuarto, "COLBUN", 0.25])
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta


def _sobrecostos(ruta):
    celdas = {(6, "A"): "encabezado"}
    fila = {
        "A": dt.datetime(2026, 8, 3), "R": 2 * 96 + 10,
        "S": 202608, "T": "SCCF", "U": "COLBUN-1", "W": 99.0,
        "AW": 1, "BC": 2, "BI": 3,          # CPF(+) = 6
        "AX": -1,                           # CPF(-) = -1 (BD, BJ vacias)
        "AY": 0.5, "BE": 0.5,               # CSF(+) = 1
        "BM": 4,                            # CTF(+) = 4
    }
    for letra, valor in fila.items():
        celdas[(7, letra)] = valor
    return _libro(ruta, celdas, hoja="SOBRECOSTOS", hojas_extra=["Otra"])


class PruebaCargaEntradas(unittest.TestCase):

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.base = Path(self._tmp.name) / "CRA 2608"
        self.base.mkdir()

    def tearDown(self):
        self._tmp.cleanup()

    def test_energia_calcula_horadia_y_periodo(self):
        ruta = _energia(self.base / "e.xlsx")

        df = cra.construir_energia(ruta, 2026, 8)

        self.assertEqual(list(df.columns), p.COLUMNAS_ENERGIA)
        self.assertEqual(len(df), 3)       # el titulo y el encabezado no entran
        self.assertEqual(df[p.CAMPO_HORADIA].tolist(), [0, 0, 23])
        # J*4 + (MINUTO+15)/15
        self.assertEqual(df[p.CAMPO_PERIODO].tolist(), [1, 4, 95])
        self.assertEqual(df[p.CAMPO_KWHD].tolist(), [100, 110, 90])
        self.assertEqual(df[p.CAMPO_KWHR].tolist(), [5, 0, 1])

    def test_energia_de_otro_mes_se_rechaza(self):
        ruta = _energia(self.base / "e.xlsx", mes=7)

        with self.assertRaises(cra.ErrorEntrada):
            cra.construir_energia(ruta, 2026, 8)

    def test_fp_y_co_copian_su_bloque(self):
        fp = _libro(self.base / "fp.xlsx", {
            (1, "A"): "BarNom", (2, "A"): "Quillota 220", (2, "B"): 3,
            (2, "C"): 1.02, (2, "D"): 1,
        })
        co = _libro(self.base / "co.xlsx", {
            (1, "A"): "nombre", (2, "A"): "RALCO_CONF", (2, "B"): 7,
            (2, "C"): 55.5, (2, "D"): 12,
        })

        df_fp = cra.construir_fp(fp)
        df_co = cra.construir_co(co)

        self.assertEqual(df_fp.values.tolist(), [["Quillota 220", 3, 1.02, 1]])
        # Orden de la hoja CO: configuracion, Dia (D), Hora (B), Costo (C)
        self.assertEqual(list(df_co.columns), list(p.COPIA_CO.values()))
        self.assertEqual(df_co.values.tolist(), [["RALCO_CONF", 12, 7, 55.5]])

    def test_sc_co_apila_co_y_sc(self):
        reporte = _reporte_cra(self.base / "r.csv")
        sobrecostos = _sobrecostos(self.base / "s.xlsx")

        avisos = []
        df = cra.construir_sc_co(reporte, sobrecostos, avisos.append)

        # SC arriba, CO abajo; solo embalses (PE-TOLPANSUR y NUEVA-7 fuera).
        self.assertEqual(list(df.columns), p.COLUMNAS_SC_CO)
        self.assertEqual(df[p.CAMPO_TIPO].tolist(), ["SCCF", "CO"])
        self.assertEqual(df[p.CAMPO_UNIDAD_SC].tolist(), ["COLBUN-1", "RAPEL-1"])
        self.assertEqual(df[p.CAMPO_CLAVE_BLOQUE].tolist(), ["3#10", "5#17"])
        self.assertTrue(any("NUEVA-7" in a and "unidad nueva" in a for a in avisos))
        self.assertFalse(any("PE-TOLPANSUR" in a and "unidad nueva" in a
                             for a in avisos))

        sc = df.iloc[0]
        self.assertEqual(sc[p.CAMPO_CPF_MAS], 6)
        self.assertEqual(sc[p.CAMPO_CPF_MENOS], -1)
        self.assertEqual(sc[p.CAMPO_CSF_MAS], 1)
        self.assertEqual(sc[p.CAMPO_CSF_MENOS], 0)
        self.assertEqual(sc[p.CAMPO_CTF_MAS], 4)

        servicios = cra.participacion_por_servicio(df)
        self.assertEqual(servicios.iloc[0].tolist(), [5.0, 1.0, 4.0])
        self.assertEqual(servicios.iloc[1].tolist(), [0.5, 0.5, 0.0])

    def test_sc_co_con_fechas_de_otro_mes_se_rechaza(self):
        with self.assertRaises(cra.ErrorEntrada):
            cra.construir_sc_co(
                _reporte_cra(self.base / "r.csv"),
                _sobrecostos(self.base / "s.xlsx"),
                periodo=(2026, 9),
            )

    def test_fd_copia_la_hoja_con_sus_encabezados(self):
        ruta = _sscc_desempeno(self.base / "SSCC_Desempeño_Agosto_2026.xlsx")

        cpf = cra.construir_fd(ruta, "fd_cpf")
        csf = cra.construir_fd(ruta, "fd_csf")

        self.assertEqual(list(cpf.columns)[:3], ["Fecha", "Hora", "Unidad"])
        self.assertEqual(len(cpf.columns), 9)            # B:J
        self.assertEqual(len(csf.columns), 7)            # B:H
        self.assertEqual(cpf.iloc[0]["Unidad"], "RAPEL_U1")
        self.assertEqual(cpf.iloc[0]["Col J"], 0.9)

    def test_prorrata_como_matriz(self):
        avisos = []
        ruta = _prorrata(self.base / "Prorrata_Retiros_2608_def.xlsx")

        df = cra.construir_matriz_prorrata(ruta, 2026, 8, avisos.append)

        self.assertEqual(list(df.columns),
                         [p.COLUMNA_PERIODO_PRORRATA, "COLBUN", "ENGIE"])
        self.assertEqual(len(df), 31 * 96)
        self.assertEqual(df.iloc[0].tolist(), [1, 0.25, 0.75])
        self.assertFalse(any("AVISO" in a for a in avisos))

        # Septiembre tiene 30 dias: 31*96 periodos sobran y se avisa.
        avisos.clear()
        cra.construir_matriz_prorrata(ruta, 2026, 9, avisos.append)
        self.assertTrue(any("sobran 96" in a for a in avisos))

    def test_hoja_no_confirmada_con_varias_hojas_se_detiene(self):
        ruta = _libro(
            self.base / "fp.xlsx", {(2, "A"): "x"}, hojas_extra=["Otra"]
        )

        with self.assertRaises(cra.ErrorEntrada) as ctx:
            cra.construir_fp(ruta)

        self.assertIn("Otra", str(ctx.exception))


REPORTE_REAL = (
    Path(__file__).resolve().parent.parent
    / "docs" / "Reporte_CRA_15min_2512_real.csv"
)


@unittest.skipUnless(REPORTE_REAL.exists(), "falta el CSV real en docs/")
class PruebaReporteCraReal(unittest.TestCase):
    """Contra el Reporte_CRA_15min_2512.csv real que entrego el usuario."""

    def test_lee_el_csv_real(self):
        df = cra.construir_sc_co_desde_reporte(REPORTE_REAL, periodo=(2025, 12))

        # 6.867 filas en el archivo; los dos parques eolicos
        # (PE-TOLPANSUR 650, PE-SANGABRIEL 531) no son embalse.
        self.assertEqual(len(df), 6867 - 650 - 531)
        self.assertFalse(df[p.CAMPO_UNIDAD_SC].str.startswith("PE-").any())
        self.assertEqual(set(df[p.CAMPO_TIPO]), {"CO"})
        self.assertEqual(set(df[p.CAMPO_CLAVE_ANIO_MES]), {2512})
        self.assertEqual(df.iloc[0][p.CAMPO_UNIDAD_SC], "CANUTILLAR-1")
        self.assertEqual(df.iloc[0][p.CAMPO_CLAVE_BLOQUE], "1#38")
        self.assertEqual(df.iloc[-1][p.CAMPO_CLAVE_BLOQUE], "16#92")

    def test_participacion_coincide_con_la_del_archivo(self):
        # El CSV trae ademas Prorrata_CPF/CSF/CTF (P, Q, R) ya sumadas:
        # tienen que dar lo mismo que CPF(+) + CPF(-), etc.
        from Script.cra.hojas_entrada import filtrar_embalses
        from Script.cra.lectura import leer_columnas

        df = cra.construir_sc_co_desde_reporte(REPORTE_REAL)
        propias = cra.participacion_por_servicio(df)

        archivo = leer_columnas(REPORTE_REAL, None, 2, ["H", "P", "Q", "R"])
        archivo = filtrar_embalses(
            archivo.rename(columns={"H": p.CAMPO_UNIDAD_SC}), "prueba"
        )

        for servicio, letra in zip(("CPF", "CSF", "CTF"), "PQR"):
            diferencia = (propias[servicio] - archivo[letra].astype(float)).abs()
            self.assertLess(diferencia.max(), 1e-6, servicio)

    def test_otro_periodo_se_rechaza(self):
        with self.assertRaises(cra.ErrorEntrada):
            cra.construir_sc_co_desde_reporte(REPORTE_REAL, periodo=(2026, 1))


class PruebaProceso(unittest.TestCase):

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.base = Path(self._tmp.name) / "CRA 2608"
        cra.crear_carpetas_caso(self.base)
        _energia(self.base / p.CARPETA_ENERGIA
                 / "Formato_Solicitud_SSAA_SSCC_Hidro_Agosto2026.xlsx")
        _libro(self.base / p.CARPETA_FP / "fp_2608_def.xlsx",
               {(2, "A"): "B1", (2, "B"): 1, (2, "C"): 1.0, (2, "D"): 1})
        _libro(self.base / p.CARPETA_CO / "cvar_cra_2608_def.xlsx",
               {(2, "A"): "C1", (2, "B"): 1, (2, "C"): 2.0, (2, "D"): 1})
        _reporte_cra(self.base / p.CARPETA_SC_CO / "Reporte_CRA_15min_2608.csv")
        _sobrecostos(self.base / p.CARPETA_SC_CO
                     / "Cálculo_SobrecostosSSCC_2608.xlsm")
        _sscc_desempeno(self.base / p.CARPETA_FD
                        / "SSCC_Desempeño_Agosto_2026_V2.xlsx")
        _prorrata(self.base / p.CARPETA_PRORRATA
                  / "Prorrata_Retiros_2608_def.xlsx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_genera_todas_y_conserva_lo_que_no_toca(self):
        salida = cra.generar_balance_cra(self.base, "2608")

        libro = load_workbook(salida, read_only=True)
        self.assertEqual(libro.sheetnames, p.ORDEN_HOJAS_SALIDA)
        libro.close()

        # Agregar una hoja ajena y regenerar solo FP: la ajena queda.
        libro = load_workbook(salida)
        libro.create_sheet("Notas")["A1"] = "no tocar"
        libro.save(salida)

        cra.generar_balance_cra(self.base, "2608", secciones=["fp"])

        libro = load_workbook(salida, read_only=True)
        self.assertEqual(libro.sheetnames[-1], "Notas")
        self.assertEqual(libro["Notas"]["A1"].value, "no tocar")
        self.assertIn(p.HOJA_ENERGIA, libro.sheetnames)
        libro.close()

    def test_dos_archivos_del_mismo_tipo_se_detiene(self):
        _libro(self.base / p.CARPETA_FP / "fp_2608_pre.xlsx", {(2, "A"): "x"})

        with self.assertRaises(cra.ErrorEntrada):
            cra.generar_balance_cra(self.base, "2608", secciones=["fp"])

    def test_estructura_marca_ok_y_pendientes(self):
        filas = cra.revisar_estructura(self.base, "2608")
        por_id = {f["id"]: f for f in filas}

        for id_ in ("energia", "fp", "co", "reporte_cra", "sobrecostos",
                    "sscc_desempeno", "prorrata"):
            self.assertEqual(por_id[id_]["estado"], "ok", id_)
        self.assertEqual(por_id["salida"]["estado"], "pendiente")

    def test_fp_y_energia_de_otro_periodo_no_se_toman(self):
        # Otro mes en el nombre: no es el archivo del periodo.
        (self.base / p.CARPETA_FP / "fp_2608_def.xlsx").rename(
            self.base / p.CARPETA_FP / "fp_2607_def.xlsx")
        (self.base / p.CARPETA_ENERGIA
         / "Formato_Solicitud_SSAA_SSCC_Hidro_Agosto2026.xlsx").rename(
            self.base / p.CARPETA_ENERGIA
            / "Formato_Solicitud_SSAA_SSCC_Hidro_Julio2026.xlsx")

        rutas = cra.resolver_rutas(self.base)
        self.assertIsNone(cra.buscar_entrada(rutas, "fp", "2608"))
        self.assertIsNone(cra.buscar_entrada(rutas, "energia", "2608"))

    def test_aamm_invalido(self):
        for malo in ("", "26", "2613", "abcd"):
            with self.assertRaises(cra.ErrorEntrada):
                cra.validar_aamm(malo)
        self.assertEqual(cra.validar_aamm("2608"), (2026, 8))


if __name__ == "__main__":
    unittest.main()
