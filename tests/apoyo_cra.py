# -*- coding: utf-8 -*-
"""Archivos sinteticos compartidos por las pruebas del CRA."""

from openpyxl import Workbook


def maestro_cra(ruta, configuraciones, empresas=None, diccionario=None):
    """
    centrales_cra.xlsx con el formato real (docs/centrales_cra_real.xlsx):
    tres hojas, titulo en B5, encabezados en la fila 8, datos desde la 9
    en la columna B.

    diccionario: [(unidad/configuracion, fd_cpf, fd_csf, fd_ctf), ...]
    """

    libro = Workbook()
    libro.remove(libro.active)
    tablas = {
        "centrales_cra": (["Configuracion"], [[c] for c in configuraciones]),
        "empresas": (["UNIDAD/CONFIGURACION", "EMPRESA"],
                     [list(f) for f in (empresas or [])]),
        "diccionario": (["UNIDAD/CONFIGURACION", "FD_CPF", "FD_CSF", "FD_CTF"],
                        [list(f) for f in (diccionario or [])]),
    }
    for hoja, (encabezados, filas) in tablas.items():
        ws = libro.create_sheet(hoja)
        ws["B5"] = f"Cuadro N° x: {hoja}"
        for j, texto in enumerate(encabezados):
            ws.cell(row=8, column=2 + j, value=texto)
        for i, fila in enumerate(filas):
            for j, valor in enumerate(fila):
                ws.cell(row=9 + i, column=2 + j, value=valor)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
    return ruta
